from __future__ import annotations

import hashlib
import json
import logging
import math
import time
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document as DocxDocument
from pgvector import Vector
from pypdf import PdfReader

from .config import settings
from .providers import embeddings_client

client = embeddings_client()


def split_text(text: str, chunk_size: int = 1200, overlap: int = 220) -> list[str]:
    if chunk_size < 1 or not 0 <= overlap < chunk_size // 2:
        raise ValueError("overlap debe ser menor que la mitad de chunk_size")
    clean = " ".join(text.replace("\x00", " ").split())
    if not clean:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + chunk_size)
        if end < len(clean):
            boundary = clean.rfind(". ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary + 1
        chunk = clean[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(clean):
            break
        start = max(end - overlap, start + 1)
    return chunks


def extract_pages(filename: str, raw: bytes) -> list[tuple[int | None, str]]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(BytesIO(raw))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if ext == ".docx":
        with ZipFile(BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("DOCX descomprimido demasiado grande")
        doc = DocxDocument(BytesIO(raw))
        return [
            (
                None,
                "\n".join(
                    [p.text for p in doc.paragraphs]
                    + [
                        " | ".join(cell.text for cell in row.cells)
                        for table in doc.tables
                        for row in table.rows
                    ]
                ),
            )
        ]
    if ext in {".txt", ".md"}:
        return [(None, raw.decode("utf-8-sig"))]
    raise ValueError("Formato no soportado. Usa PDF, DOCX, TXT o MD.")


def embedding_collection():
    if settings.demo_mode:
        settings.guard_demo()
        return f"demo:deterministic-v1:{settings.vector_dimensions}:fixture-only"
    revision = settings.embedding_revision
    if settings.embedding_provider == "ollama":
        revision = client.digest(settings.embedding_model) + ":retrieval-prefix-v1"
    if settings.embedding_provider == "gemini":
        revision += ":retrieval-prefix-v1"
    return f"{settings.embedding_provider}:{settings.embedding_model}:{settings.vector_dimensions}:{revision}"


def embed_texts(texts: list[str], *, query=False) -> list[list[float]]:
    before = embedding_collection()
    inputs = texts
    if (
        settings.embedding_provider == "ollama"
        and settings.embedding_model.split(":")[0] == "embeddinggemma"
    ) or settings.embedding_provider == "gemini":
        prefix = "task: search result | query: " if query else "title: none | text: "
        inputs = [prefix + text for text in texts]
    result = client.embeddings.create(
        model=settings.embedding_model, input=inputs, dimensions=settings.vector_dimensions
    )
    ordered = sorted(result.data, key=lambda item: item.index)
    if [item.index for item in ordered] != list(range(len(texts))) or any(
        len(item.embedding) != settings.vector_dimensions
        or not all(math.isfinite(v) for v in item.embedding)
        or not any(item.embedding)
        for item in ordered
    ):
        raise RuntimeError("Respuesta de embeddings incompatible con la colección configurada")
    if before != embedding_collection():
        raise RuntimeError("El modelo cambió durante el embedding; vuelve a intentarlo")
    return [item.embedding for item in ordered]


def embed_query(text: str) -> list[float]:
    return embed_texts([text], query=True)[0]


def retrieve(
    conn, workspace_id: str, query: str, limit: int = 6, *, diagnostics=None, lexical_query=None
) -> list[dict]:
    """Workspace-scoped exact vector candidates + Spanish FTS, fused with RRF.

    Exact scoped distance avoids HNSW post-filter starvation across tenants.
    This deliberately favors recall for the current small corpus; benchmark before scaling.
    Cosine score remains separate from the fusion score (neither is confidence).
    """
    if not 1 <= limit <= 20:
        raise ValueError("limit debe estar entre 1 y 20")
    started = time.perf_counter()
    collection = embedding_collection()
    vector = Vector(embed_query(query))
    if collection != embedding_collection():
        raise RuntimeError("La colección cambió durante la consulta")
    candidate_limit = max(24, limit * 4)
    columns = "c.id,c.document_id,d.filename,c.page_number,c.content"
    filters = """c.workspace_id=%s AND d.workspace_id=%s
                 AND d.status='ready' AND d.embedding_collection=%s
                 AND c.embedding_collection=d.embedding_collection"""
    args = (workspace_id, workspace_id, collection)
    semantic = conn.execute(
        f"""WITH scoped AS MATERIALIZED (
            SELECT {columns},c.embedding FROM document_chunks c
            JOIN documents d ON d.id=c.document_id WHERE {filters}
        ) SELECT id,document_id,filename,page_number,content,
                 1-(embedding <=> %s) AS score
          FROM scoped ORDER BY embedding <=> %s,id LIMIT %s""",
        (*args, vector, vector, candidate_limit),
    ).fetchall()
    lexical = conn.execute(
        f"""WITH scoped AS MATERIALIZED (
            SELECT {columns},c.embedding,c.search_vector FROM document_chunks c
            JOIN documents d ON d.id=c.document_id WHERE {filters}
        ) SELECT id,document_id,filename,page_number,content,
                 1-(embedding <=> %s) AS score,ts_rank_cd(search_vector,q) AS lexical_score
          FROM scoped,websearch_to_tsquery('spanish',%s) q
          WHERE search_vector @@ q ORDER BY lexical_score DESC,id LIMIT %s""",
        (*args, vector, lexical_query if lexical_query is not None else query, candidate_limit),
    ).fetchall()
    semantic = [r for r in semantic if math.isfinite(r[5]) and r[5] >= settings.retrieval_min_score]
    candidates = {}
    for kind, rows in (("vector", semantic), ("lexical", lexical)):
        for rank, row in enumerate(rows, 1):
            item = candidates.setdefault(
                row[0],
                {
                    "chunk_id": row[0],
                    "document_id": str(row[1]),
                    "filename": row[2],
                    "page": row[3],
                    "excerpt": row[4],
                    "content": row[4],
                    "score": float(row[5]) if math.isfinite(row[5]) else 0.0,
                    "rrf_score": 0.0,
                    "vector_rank": None,
                    "lexical_rank": None,
                    "lexical_score": 0.0,
                },
            )
            item["rrf_score"] += 1 / (60 + rank)
            item[kind + "_rank"] = rank
            if kind == "lexical":
                item["lexical_score"] = float(row[6])
    selected, seen = [], set()
    for item in sorted(candidates.values(), key=lambda s: (-s["rrf_score"], s["chunk_id"])):
        fingerprint = hashlib.sha256(" ".join(item["content"].casefold().split()).encode()).digest()
        if fingerprint in seen:
            continue
        seen.add(fingerprint)
        selected.append(item)
        if len(selected) == limit:
            break
    meta = {
        "strategy": "hybrid_rrf_v1",
        "embedding_collection": collection,
        "vector_candidates": len(semantic),
        "lexical_candidates": len(lexical),
        "unique_candidates": len(candidates),
        "selected": [
            {
                k: s[k]
                for k in (
                    "chunk_id",
                    "document_id",
                    "score",
                    "rrf_score",
                    "vector_rank",
                    "lexical_rank",
                    "lexical_score",
                )
            }
            for s in selected
        ],
        "duration_ms": round((time.perf_counter() - started) * 1000, 2),
        "abstention_reason": None if selected else "no_retrieval_evidence",
    }
    if diagnostics is not None:
        diagnostics.update(meta)
    logging.getLogger(__name__).info("retrieval %s", json.dumps(meta))
    return selected
